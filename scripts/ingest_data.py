from __future__ import annotations

import argparse
import json
from typing import Any

import psycopg
import requests
from openpyxl import load_workbook

from db_utils import (
    PROJECT_ROOT,
    chunk_text,
    find_data_root,
    formation_info,
    format_vector,
    infer_year_from_label,
    infer_year_from_semester,
    load_env_file,
    normalize_formation_name,
    normalize_group_name,
    normalize_text,
    parse_decimal,
    parse_french_date_range,
    semester_year_from_path,
)


def simplify_sheet(value: str) -> str:
    return " ".join(normalize_text(value).lower().replace("-", " ").split())


class DataIngestor:
    def __init__(
        self,
        database_url: str,
        ollama_base_url: str,
        embedding_model: str,
        embedding_dimension: int,
        academic_start_year: int,
        skip_embeddings: bool = False,
        reset: bool = False,
    ) -> None:
        self.database_url = database_url
        self.ollama_base_url = ollama_base_url.rstrip("/")
        self.embedding_model = embedding_model
        self.embedding_dimension = embedding_dimension
        self.academic_start_year = academic_start_year
        self.skip_embeddings = skip_embeddings
        self.reset = reset
        self.data_root = find_data_root()
        self.conn = psycopg.connect(self.database_url)
        self.conn.autocommit = False

        self.formation_cache: dict[str, int] = {}
        self.niveau_cache: dict[tuple[int, int], int] = {}
        self.semestre_cache: dict[tuple[int, int], int] = {}
        self.group_cache: dict[tuple[int, int, str], int] = {}
        self.ue_cache: dict[tuple[int, str, str], int] = {}

    def close(self) -> None:
        self.conn.close()

    def run(self) -> None:
        try:
            with self.conn.cursor() as cur:
                if self.reset:
                    self._reset_tables(cur)
                self._seed_formations(cur)
                self._seed_academic_structure(cur)
                self._ingest_calendar(cur)
                self._ingest_descriptions(cur)
                self._ingest_groups(cur)
                self._ingest_emplois(cur)
                self._ingest_plan_etudes(cur)
            self.conn.commit()
        except Exception:
            self.conn.rollback()
            raise

    def _reset_tables(self, cur: psycopg.Cursor[Any]) -> None:
        cur.execute(
            """
            TRUNCATE TABLE
                document_chunks,
                documents,
                emplois_temps,
                etudiants,
                matieres,
                unites_enseignement,
                evenements_calendrier,
                groupes,
                semestres,
                niveaux,
                formations
            RESTART IDENTITY CASCADE
            """
        )

    def _seed_formations(self, cur: psycopg.Cursor[Any]) -> None:
        for canonical in ("informatique", "infotronique", "mecatronique", "industrielle"):
            info = formation_info(canonical)
            cur.execute(
                """
                INSERT INTO formations (code, nom, description_courte)
                VALUES (%s, %s, %s)
                ON CONFLICT (code)
                DO UPDATE SET nom = EXCLUDED.nom
                RETURNING id
                """,
                (info.code, info.nom, None),
            )
            self.formation_cache[canonical] = cur.fetchone()[0]

    def _seed_academic_structure(self, cur: psycopg.Cursor[Any]) -> None:
        for canonical in ("informatique", "infotronique", "mecatronique", "industrielle"):
            formation_id = self._formation_id(cur, canonical)
            for year in (1, 2, 3):
                self._niveau_id(cur, formation_id, year)
            for semester in (1, 2, 3, 4, 5, 6):
                self._semestre_id(cur, formation_id, semester)

    def _formation_id(self, cur: psycopg.Cursor[Any], formation_raw: str) -> int:
        canonical = normalize_formation_name(formation_raw)
        cached = self.formation_cache.get(canonical)
        if cached:
            return cached
        info = formation_info(canonical)
        cur.execute("SELECT id FROM formations WHERE code = %s", (info.code,))
        row = cur.fetchone()
        if row:
            self.formation_cache[canonical] = row[0]
            return row[0]
        cur.execute(
            "INSERT INTO formations (code, nom) VALUES (%s, %s) RETURNING id",
            (info.code, info.nom),
        )
        formation_id = cur.fetchone()[0]
        self.formation_cache[canonical] = formation_id
        return formation_id

    def _niveau_id(self, cur: psycopg.Cursor[Any], formation_id: int, year: int) -> int:
        key = (formation_id, year)
        cached = self.niveau_cache.get(key)
        if cached:
            return cached
        cur.execute(
            """
            INSERT INTO niveaux (formation_id, annee)
            VALUES (%s, %s)
            ON CONFLICT (formation_id, annee)
            DO UPDATE SET annee = EXCLUDED.annee
            RETURNING id
            """,
            (formation_id, year),
        )
        niveau_id = cur.fetchone()[0]
        self.niveau_cache[key] = niveau_id
        return niveau_id

    def _semestre_id(self, cur: psycopg.Cursor[Any], formation_id: int, semester: int) -> int:
        year = infer_year_from_semester(semester)
        niveau_id = self._niveau_id(cur, formation_id, year)
        key = (niveau_id, semester)
        cached = self.semestre_cache.get(key)
        if cached:
            return cached
        cur.execute(
            """
            INSERT INTO semestres (niveau_id, numero)
            VALUES (%s, %s)
            ON CONFLICT (niveau_id, numero)
            DO UPDATE SET numero = EXCLUDED.numero
            RETURNING id
            """,
            (niveau_id, semester),
        )
        semestre_id = cur.fetchone()[0]
        self.semestre_cache[key] = semestre_id
        return semestre_id

    def _group_id(self, cur: psycopg.Cursor[Any], formation: str, year: int, group_name: str) -> int:
        formation_id = self._formation_id(cur, formation)
        self._niveau_id(cur, formation_id, year)
        key = (formation_id, year, group_name)
        cached = self.group_cache.get(key)
        if cached:
            return cached
        cur.execute(
            """
            INSERT INTO groupes (formation_id, annee, nom_groupe)
            VALUES (%s, %s, %s)
            ON CONFLICT (formation_id, annee, nom_groupe)
            DO UPDATE SET nom_groupe = EXCLUDED.nom_groupe
            RETURNING id
            """,
            (formation_id, year, group_name),
        )
        group_id = cur.fetchone()[0]
        self.group_cache[key] = group_id
        return group_id

    def _upsert_etudiant(
        self,
        cur: psycopg.Cursor[Any],
        groupe_id: int,
        nom_complet: str,
        email: str | None,
    ) -> None:
        cur.execute(
            """
            SELECT id
            FROM etudiants
            WHERE groupe_id = %s
              AND COALESCE(email, '') = COALESCE(%s, '')
              AND nom_complet = %s
            """,
            (groupe_id, email, nom_complet),
        )
        if cur.fetchone():
            return

        cur.execute(
            """
            INSERT INTO etudiants (groupe_id, nom_complet, email)
            VALUES (%s, %s, %s)
            """,
            (groupe_id, nom_complet, email),
        )

    def _insert_document(
        self,
        cur: psycopg.Cursor[Any],
        type_document: str,
        titre: str,
        formation_id: int | None,
        niveau_id: int | None,
        semestre_id: int | None,
        source_path: str,
        contenu_texte: str,
        metadata: dict[str, Any],
    ) -> int:
        cur.execute(
            """
            SELECT id
            FROM documents
            WHERE type_document = %s
              AND source_path = %s
            """,
            (type_document, source_path),
        )
        row = cur.fetchone()
        if row:
            document_id = row[0]
            cur.execute(
                """
                UPDATE documents
                SET titre = %s,
                    formation_id = %s,
                    niveau_id = %s,
                    semestre_id = %s,
                    contenu_texte = %s,
                    metadata = %s
                WHERE id = %s
                """,
                (
                    titre,
                    formation_id,
                    niveau_id,
                    semestre_id,
                    contenu_texte,
                    json.dumps(metadata, ensure_ascii=False),
                    document_id,
                ),
            )
            return document_id
        cur.execute(
            """
            INSERT INTO documents (
                type_document, titre, formation_id, niveau_id, semestre_id,
                source_path, contenu_texte, metadata
            )
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
            RETURNING id
            """,
            (
                type_document,
                titre,
                formation_id,
                niveau_id,
                semestre_id,
                source_path,
                contenu_texte,
                json.dumps(metadata, ensure_ascii=False),
            ),
        )
        return cur.fetchone()[0]

    def _embed_texts(self, texts: list[str]) -> list[list[float]]:
        if self.skip_embeddings:
            return [[] for _ in texts]

        response = requests.post(
            f"{self.ollama_base_url}/api/embed",
            json={
                "model": self.embedding_model,
                "input": texts,
                "dimensions": self.embedding_dimension,
            },
            timeout=120,
        )
        response.raise_for_status()
        payload = response.json()
        embeddings = payload.get("embeddings", [])
        if len(embeddings) != len(texts):
            raise RuntimeError("Le nombre d'embeddings retournes par Ollama est incorrect.")
        return embeddings

    def _upsert_chunk(
        self,
        cur: psycopg.Cursor[Any],
        document_id: int,
        chunk_index: int,
        contenu: str,
        embedding: list[float] | None,
        metadata: dict[str, Any],
    ) -> None:
        vector_value = format_vector(embedding) if embedding else None
        cur.execute(
            """
            INSERT INTO document_chunks (document_id, chunk_index, contenu, embedding, metadata)
            VALUES (%s, %s, %s, CAST(%s AS vector), %s)
            ON CONFLICT (document_id, chunk_index)
            DO UPDATE SET
                contenu = EXCLUDED.contenu,
                embedding = EXCLUDED.embedding,
                metadata = EXCLUDED.metadata
            """,
            (
                document_id,
                chunk_index,
                contenu,
                vector_value,
                json.dumps(metadata, ensure_ascii=False),
            ),
        )

    def _ingest_calendar(self, cur: psycopg.Cursor[Any]) -> None:
        calendar_path = self.data_root / "Calendrier" / "calendrier.json"
        if not calendar_path.exists():
            return

        entries = json.loads(calendar_path.read_text(encoding="utf-8"))
        for entry in entries:
            label = normalize_text(entry.get("date"))
            event = normalize_text(entry.get("evenement"))
            date_debut, date_fin = parse_french_date_range(label, self.academic_start_year)
            cur.execute(
                """
                SELECT id
                FROM evenements_calendrier
                WHERE COALESCE(date_debut, DATE '1900-01-01') = COALESCE(%s, DATE '1900-01-01')
                  AND COALESCE(date_fin, DATE '1900-01-01') = COALESCE(%s, DATE '1900-01-01')
                  AND libelle = %s
                """,
                (date_debut, date_fin, event),
            )
            if cur.fetchone():
                continue
            cur.execute(
                """
                INSERT INTO evenements_calendrier (date_debut, date_fin, libelle, source_label)
                VALUES (%s, %s, %s, %s)
                """,
                (date_debut, date_fin, event, label),
            )

    def _ingest_descriptions(self, cur: psycopg.Cursor[Any]) -> None:
        descriptions_dir = self.data_root / "description_formation"
        if not descriptions_dir.exists():
            return

        for path in sorted(descriptions_dir.glob("*.txt")):
            formation = normalize_formation_name(path.stem)
            formation_id = self._formation_id(cur, formation)
            title = f"Description de la formation {formation_info(formation).nom}"
            content = normalize_text(path.read_text(encoding="utf-8"))
            short_description = content[:500].strip() or None
            cur.execute(
                """
                UPDATE formations
                SET description_courte = %s
                WHERE id = %s
                """,
                (short_description, formation_id),
            )
            document_id = self._insert_document(
                cur=cur,
                type_document="description",
                titre=title,
                formation_id=formation_id,
                niveau_id=None,
                semestre_id=None,
                source_path=str(path.relative_to(PROJECT_ROOT)),
                contenu_texte=content,
                metadata={"formation": formation, "source_type": "txt"},
            )
            chunks = chunk_text(content)
            if not chunks:
                continue
            cur.execute("DELETE FROM document_chunks WHERE document_id = %s", (document_id,))
            embeddings = self._embed_texts(chunks)
            for index, chunk in enumerate(chunks):
                embedding = embeddings[index] if not self.skip_embeddings else None
                self._upsert_chunk(
                    cur,
                    document_id=document_id,
                    chunk_index=index,
                    contenu=chunk,
                    embedding=embedding,
                    metadata={"formation": formation, "source_type": "txt"},
                )

    def _ingest_groups(self, cur: psycopg.Cursor[Any]) -> None:
        groups_root = self.data_root / "Liste_groupes"
        if not groups_root.exists():
            return

        for workbook_path in sorted(groups_root.glob("*/*/*.xlsx")):
            if "email_groupe+email_perso" not in str(workbook_path):
                continue
            formation = normalize_formation_name(workbook_path.parts[-3])
            year = int(workbook_path.stem[0])
            workbook = load_workbook(workbook_path, read_only=True, data_only=True)
            try:
                for sheet_name in workbook.sheetnames:
                    sheet_clean = normalize_text(sheet_name)
                    if not sheet_clean or simplify_sheet(sheet_clean) == "feuille 1":
                        continue
                    group_name = normalize_group_name(sheet_clean, formation, year_hint=year)
                    group_id = self._group_id(cur, formation, year, group_name)
                    sheet = workbook[sheet_name]
                    for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                        if row_index == 1:
                            continue
                        group_email = normalize_text(row[0]) if len(row) > 0 else ""
                        member_email = normalize_text(row[1]) if len(row) > 1 else ""
                        member_name = normalize_text(row[2]) if len(row) > 2 else ""
                        if not group_email and not member_email and not member_name:
                            continue
                        if not member_name:
                            continue
                        self._upsert_etudiant(
                            cur,
                            groupe_id=group_id,
                            nom_complet=member_name,
                            email=member_email or None,
                        )
            finally:
                workbook.close()

    def _ingest_emplois(self, cur: psycopg.Cursor[Any]) -> None:
        emplois_root = self.data_root / "emplois"
        if not emplois_root.exists():
            return

        for path in sorted(emplois_root.rglob("*.json")):
            rows = json.loads(path.read_text(encoding="utf-8"))
            for row in rows:
                formation = normalize_formation_name(row.get("filiere") or path.parent.name)
                local_semester = int(row["semestre"])
                year = infer_year_from_label(str(row["groupe"]))
                semester = ((year - 1) * 2) + local_semester
                group_name = normalize_group_name(row["groupe"], formation, year_hint=year)
                group_id = self._group_id(cur, formation, year, group_name)

                cur.execute(
                    """
                    SELECT id
                    FROM emplois_temps
                    WHERE groupe_id = %s
                      AND semestre = %s
                      AND jour = %s
                      AND heure_debut = %s
                      AND heure_fin = %s
                      AND matiere = %s
                      AND COALESCE(enseignant, '') = COALESCE(%s, '')
                      AND COALESCE(salle, '') = COALESCE(%s, '')
                      AND COALESCE(type_seance, '') = COALESCE(%s, '')
                    """,
                    (
                        group_id,
                        semester,
                        normalize_text(row["jour"]).lower(),
                        row["heure_debut"],
                        row["heure_fin"],
                        normalize_text(row["matiere"]),
                        normalize_text(row.get("enseignant")),
                        normalize_text(row.get("salle")),
                        normalize_text(row.get("type_seance")),
                    ),
                )
                if cur.fetchone():
                    continue

                cur.execute(
                    """
                    INSERT INTO emplois_temps (
                        groupe_id, semestre, jour, heure_debut, heure_fin,
                        matiere, enseignant, salle, type_seance
                    )
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
                    """,
                    (
                        group_id,
                        semester,
                        normalize_text(row["jour"]).lower(),
                        row["heure_debut"],
                        row["heure_fin"],
                        normalize_text(row["matiere"]),
                        normalize_text(row.get("enseignant")),
                        normalize_text(row.get("salle")),
                        normalize_text(row.get("type_seance")),
                    ),
                )

    def _ingest_plan_etudes(self, cur: psycopg.Cursor[Any]) -> None:
        plan_root = self.data_root / "Plan_etudes"
        if not plan_root.exists():
            return

        for path in sorted(plan_root.rglob("*.xlsx")):
            formation = normalize_formation_name(
                path.parts[-2]
                if path.parts[-2] in {"industrielle", "informatique", "infotronique", "mecatronique"}
                else path.parts[-3]
            )
            if formation == "industrielle":
                continue

            _, semester = semester_year_from_path(path)
            formation_id = self._formation_id(cur, formation)
            semestre_id = self._semestre_id(cur, formation_id, semester)

            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                sheet = workbook[workbook.sheetnames[0]]
                current_ue_id: int | None = None
                for row in sheet.iter_rows(values_only=True):
                    row_values = list(row[:10]) if row else []
                    if len(row_values) < 10:
                        row_values.extend([None] * (10 - len(row_values)))
                    values = [normalize_text(cell) for cell in row_values]
                    if not any(values):
                        continue
                    col_a = values[0]
                    col_b = values[1]
                    col_a_normalized = normalize_text(col_a).lower()
                    col_b_normalized = normalize_text(col_b)
                    if col_a_normalized.startswith("unité d") or col_a_normalized.startswith("sous-totaux"):
                        continue
                    if col_b_normalized == "Eléments d’Enseignement (EE)":
                        continue
                    if not col_b:
                        continue

                    if col_a.startswith("UE"):
                        current_ue_id = self._upsert_ue(
                            cur=cur,
                            semestre_id=semestre_id,
                            code_ue=col_a,
                            nom_ue=col_a,
                            coefficient_ue=parse_decimal(values[7]),
                            credits_ue=None,
                        )

                    if current_ue_id is None:
                        continue

                    self._upsert_matiere(
                        cur=cur,
                        ue_id=current_ue_id,
                        nom=col_b,
                        ci=parse_decimal(values[2]),
                        td=parse_decimal(values[3]),
                        tp=parse_decimal(values[4]),
                        total=parse_decimal(values[5]),
                        coefficient=parse_decimal(values[6]),
                        evaluation_cc=values[8] or None,
                        evaluation_examen=values[9] or None,
                    )
            finally:
                workbook.close()

    def _upsert_ue(
        self,
        cur: psycopg.Cursor[Any],
        semestre_id: int,
        code_ue: str,
        nom_ue: str,
        coefficient_ue: float | None,
        credits_ue: float | None,
    ) -> int:
        key = (semestre_id, code_ue, nom_ue)
        cached = self.ue_cache.get(key)
        if cached:
            return cached
        cur.execute(
            """
            SELECT id
            FROM unites_enseignement
            WHERE semestre_id = %s
              AND COALESCE(code_ue, '') = COALESCE(%s, '')
            """,
            (semestre_id, code_ue),
        )
        row = cur.fetchone()
        if row:
            ue_id = row[0]
            cur.execute(
                """
                UPDATE unites_enseignement
                SET nom_ue = %s,
                    coefficient_ue = %s,
                    credits_ue = %s
                WHERE id = %s
                """,
                (nom_ue, coefficient_ue, credits_ue, ue_id),
            )
            self.ue_cache[key] = ue_id
            return ue_id
        cur.execute(
            """
            INSERT INTO unites_enseignement (
                semestre_id, code_ue, nom_ue, coefficient_ue, credits_ue
            )
            VALUES (%s, %s, %s, %s, %s)
            RETURNING id
            """,
            (semestre_id, code_ue, nom_ue, coefficient_ue, credits_ue),
        )
        ue_id = cur.fetchone()[0]
        self.ue_cache[key] = ue_id
        return ue_id

    def _upsert_matiere(
        self,
        cur: psycopg.Cursor[Any],
        ue_id: int,
        nom: str,
        ci: float | None,
        td: float | None,
        tp: float | None,
        total: float | None,
        coefficient: float | None,
        evaluation_cc: str | None,
        evaluation_examen: str | None,
    ) -> None:
        cur.execute(
            """
            SELECT id
            FROM matieres
            WHERE ue_id = %s
              AND nom = %s
              AND COALESCE(ci, -1) = COALESCE(%s, -1)
              AND COALESCE(td, -1) = COALESCE(%s, -1)
              AND COALESCE(tp, -1) = COALESCE(%s, -1)
              AND COALESCE(total, -1) = COALESCE(%s, -1)
              AND COALESCE(coefficient, -1) = COALESCE(%s, -1)
            """,
            (ue_id, nom, ci, td, tp, total, coefficient),
        )
        if cur.fetchone():
            return
        cur.execute(
            """
            INSERT INTO matieres (
                ue_id, nom, ci, td, tp, total, coefficient, evaluation_cc, evaluation_examen
            )
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
            """,
            (ue_id, nom, ci, td, tp, total, coefficient, evaluation_cc, evaluation_examen),
        )


def build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Insere les donnees universitaires dans PostgreSQL.")
    parser.add_argument(
        "--academic-start-year",
        type=int,
        default=2025,
        help="Annee de debut de l'annee universitaire. Ex: 2025 pour 2025-2026.",
    )
    parser.add_argument(
        "--skip-embeddings",
        action="store_true",
        help="Insere les documents sans appeler Ollama pour generer les embeddings.",
    )
    parser.add_argument(
        "--reset",
        action="store_true",
        help="Vide les tables cibles avant l'insertion.",
    )
    return parser


def main() -> None:
    args = build_argument_parser().parse_args()
    env = load_env_file()

    database_url = env.get("DATABASE_URL")
    if not database_url:
        raise RuntimeError("DATABASE_URL est introuvable dans .env")

    ingestor = DataIngestor(
        database_url=database_url,
        ollama_base_url=env.get("OLLAMA_BASE_URL", "http://localhost:11434"),
        embedding_model=env.get("OLLAMA_EMBEDDING_MODEL", "embeddinggemma"),
        embedding_dimension=int(env.get("PGVECTOR_DIMENSION", "768")),
        academic_start_year=args.academic_start_year,
        skip_embeddings=args.skip_embeddings,
        reset=args.reset,
    )
    try:
        ingestor.run()
    finally:
        ingestor.close()

    print("Insertion terminee avec succes.")


if __name__ == "__main__":
    main()
